"""Insert a one-line introductory sentence before every figure that
currently lacks one (image paragraph preceded by a soft heading or
another figure's caption).

Placement: the new paragraph is inserted immediately before the image
paragraph, so it reads "intro → figure → caption".

Each figure is identified by a unique substring from its caption.
"""

from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
DOC = ROOT / "project_report_v2.docx"

# (caption_prefix, intro_sentence) — caption_prefix must be a unique
# substring from the existing figure caption so we can locate it.
INTROS = [
    # §3.3 Data Splitting
    ("Figure 5.1: Data Splitting Visualization — Leave-Last-Out",
     "The figure below visualises the LLO split: for every user, the most "
     "recent interaction is held out as the test item and everything earlier "
     "is used for training."),

    # §4.2.1 MF-BPR training
    ("Figure 6: MF-BPR training process",
     "The figure below walks through one training step end-to-end."),

    # Per-lesson illustrations (Figure L1..L8)
    ("Figure L1: Pointwise vs pairwise",
     "Before looking at the numbers, it helps to see why the two training "
     "objectives diverge so sharply on identical data."),
    ("Figure L2: At 5.2% overlap the bridge population is tiny",
     "The bridge population at 5.2% overlap is thin — the figure below "
     "makes the scale explicit."),
    ("Figure L3: Moving from 5.2% → 100% overlap",
     "To visualise what the overlap filter buys, we compare the Lesson 2 "
     "mixed cohort against the Lesson 3 all-overlap cohort."),
    ("Figure L4: % behind LightGCN, by model",
     "Figure L4 visualises each CDR family's gap to LightGCN as we tighten "
     "the source-domain minimum from ≥5 to ≥10 movies."),
    ("Figure L5: Popularity-based trim",
     "The catalog-sharpening step aggressively trims long-tail movie items. "
     "Figure L5 shows the trim magnitude and how each model responds."),
    ("Figure L6: User-split cold-start",
     "The cold-start regime is where single-domain baselines are most "
     "exposed. Figure L6 lays out the user-split protocol and the collapse "
     "it produces."),
    ("Figure L7: Overall parity panel vs niche subgroup panel",
     "Collaborative filtering cannot reach a user with no interaction "
     "signal, but a content-aware SBERT bridge can. Figure L7 compares "
     "overall performance against the niche-item subgroup."),
    ("Figure L8: Co-occurrence rerank pipeline",
     "Figure L8 traces the rerank pipeline and plots per-model before/after "
     "deltas across four deployments."),

    # Per-lesson findings graphs (Figure L#b)
    ("Figure L1b: MF-BPR vs MF-Explicit on Recall@10 and NDCG@10",
     "Figure L1b places the two objectives side by side on Recall@10 and "
     "NDCG@10."),
    ("Figure L2b: Recall@10 on the mixed cohort",
     "Figure L2b plots Recall@10 for every model on the mixed cohort."),
    ("Figure L3b: L2 vs L3 Recall@10 per model",
     "Figure L3b shows the L2 → L3 lift per model."),
    ("Figure L4b: L3 vs L4 gap",
     "Figure L4b plots each model's gap to LightGCN across L3 and L4."),
    ("Figure L5b: L4 vs L5 Recall@10",
     "Figure L5b plots L4 vs L5 Recall@10 per model."),
    ("Figure L6b: Cold-start Recall@10 across six models",
     "Figure L6b plots cold-start Recall@10 across six models."),
    ("Figure L7b: Overall vs niche Recall@10",
     "Figure L7b puts the overall panel next to the niche subgroup panel."),
    ("Figure L8b: Base vs rerank Recall@10 across four deployments",
     "Figure L8b plots base vs rerank Recall@10 across four deployments."),

    # Back-to-back figure pairs (second figure gets its own lead sentence)
    ("Figure 20: Lesson 1 — MF-BPR vs MF-Explicit comparison",
     "Figure 20 plots the same comparison on the L1 numerics."),
    ("Figure 23: Lesson 3 — 100% overlap users",
     "Figure 23 breaks the same split down per model on the "
     "100%-overlap cohort."),
    ("Figure 27: Lesson 6 — Cold-start: CDR vs single-domain",
     "Figure 27 expands the comparison to include popularity and "
     "content-based baselines."),
    ("Figure 29: Lesson 7 — Subgroup analysis: SBERT dominates niche items",
     "Figure 29 breaks the overall result into popularity subgroups."),
    ("Figure 40 — Nine recommendation rows",
     "Figure 40 below visualises the nine-row assembly, mapping each "
     "frontend row to the lesson that motivated it."),
]


def _find_caption_para(doc, prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    return None


def _find_image_before(doc, caption_para):
    """Walk paragraphs backward from the caption and return the nearest one
    that contains a drawing/inline image."""
    ns = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
    paragraphs = doc.paragraphs
    idx = None
    for i, p in enumerate(paragraphs):
        if p._p is caption_para._p:
            idx = i
            break
    if idx is None:
        return None
    for j in range(idx - 1, max(idx - 5, -1), -1):
        p = paragraphs[j]
        if any(r._element.findall(f".//{ns}inline") for r in p.runs):
            return p
        if any(r._element.findall(f".//{ns}anchor") for r in p.runs):
            return p
    return None


def _new_body_paragraph(text, template_paragraph):
    """Create a new body paragraph that copies the template's pPr (style,
    spacing) but carries a fresh run with just `text` in Arial 10pt."""
    new_p = deepcopy(template_paragraph._p)
    # Drop all runs from clone, keep pPr.
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    for r in list(new_p.findall(f"{ns}r")):
        new_p.remove(r)

    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rFonts = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "cs"):
        rFonts.set(qn(f"w:{attr}"), "Arial")
    rPr.append(rFonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "20")  # 10pt = half-points 20
    rPr.append(sz)
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), "20")
    rPr.append(szCs)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    new_p.append(r)
    return new_p


def _insert_before(existing_para, new_p_elem):
    parent = existing_para._p.getparent()
    parent.insert(list(parent).index(existing_para._p), new_p_elem)


def main():
    doc = Document(str(DOC))
    applied = 0
    missing = []
    # Pick a reliable template paragraph (a known normal body paragraph).
    template = None
    for p in doc.paragraphs:
        if p.style.name == "normal" and p.text.strip():
            template = p
            break

    for prefix, intro in INTROS:
        cap = _find_caption_para(doc, prefix)
        if cap is None:
            missing.append(prefix)
            continue
        img = _find_image_before(doc, cap)
        if img is None:
            missing.append(f"(no image before) {prefix}")
            continue
        new_elem = _new_body_paragraph(intro, template)
        _insert_before(img, new_elem)
        applied += 1

    doc.save(str(DOC))
    print(f"Inserted {applied} figure intros")
    for m in missing:
        print(f"  ✗ {m}")


if __name__ == "__main__":
    main()
