"""Rewrite §4.3.1 CMF: Collective Matrix Factorization.

Actions
-------
1. Under "CMF Architecture" — expand the single-paragraph description into
   a deeper walkthrough of the three-matrix design (shared U, per-domain
   V_m and V_g) and why the shared user matrix *is* the transfer
   mechanism.
2. Under "CMF Training" — add body content (initialisation, sampling,
   per-domain BPR streams, α weighting, regularisation, SGD), and move
   Equation 4 + its "where …" symbol explanation *into* this subsection.
   Figure 13 (the mechanism diagram) is repositioned here, right after
   the loss discussion.
3. Under "CMF Inference" — keep the existing body paragraph and append
   a new illustration (fig13b_cmf_inference.png + caption).

The script operates by locating the four H4 anchors (Architecture,
Training, Inference, Limitations) and surgically replacing / augmenting
content between them.
"""

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
DOC = ROOT / "project_report_v2.docx"
FIGS = ROOT / "report_figures_v3"

IMAGE_WIDTH_IN = 5.8
FORMULA_WIDTH_IN = 5.8


# ---- helpers --------------------------------------------------------

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _find_heading(doc, text):
    for p in doc.paragraphs:
        if p.style.name.startswith("Heading") and p.text.strip() == text:
            return p
    raise ValueError(f"not found: {text!r}")


def _next_paragraph_element(elem):
    """Given a w:p element, return the next sibling w:p (skipping tables),
    or None."""
    sib = elem.getnext()
    while sib is not None and sib.tag != f"{W}p":
        sib = sib.getnext()
    return sib


def _iter_between(start_elem, stop_elem):
    """Iterate paragraph elements strictly between start_elem and stop_elem."""
    cur = start_elem.getnext()
    while cur is not None and cur is not stop_elem:
        nxt = cur.getnext()
        yield cur
        cur = nxt


def _delete_between(doc, start_heading, stop_heading):
    """Remove every element (paragraphs + tables) strictly between the two
    heading paragraphs."""
    start_elem = start_heading._p
    stop_elem = stop_heading._p
    parent = start_elem.getparent()
    cur = start_elem.getnext()
    while cur is not None and cur is not stop_elem:
        nxt = cur.getnext()
        parent.remove(cur)
        cur = nxt


def _blank_p_after(anchor, alignment=None, space_before=None, space_after=None,
                   line_spacing=1.35, style_name="normal"):
    """Insert a brand-new blank w:p after anchor's w:p, return Paragraph.

    A fresh pPr is built with the requested style so we don't inherit
    Heading 4 from an H4 anchor.
    """
    from docx.text.paragraph import Paragraph
    new_p = OxmlElement("w:p")
    pPr = OxmlElement("w:pPr")
    pStyle = OxmlElement("w:pStyle")
    pStyle.set(qn("w:val"), style_name)
    pPr.append(pStyle)
    new_p.append(pPr)
    anchor._p.addnext(new_p)
    para = Paragraph(new_p, anchor._parent)
    if alignment is not None:
        para.alignment = alignment
    if space_before is not None:
        para.paragraph_format.space_before = space_before
    if space_after is not None:
        para.paragraph_format.space_after = space_after
    para.paragraph_format.line_spacing = line_spacing
    return para


def _add_arial(para, text, size_pt=10, bold=False, italic=False):
    run = para.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    # Set cs/hAnsi slots.
    rPr = run._element.find(qn("w:rPr"))
    if rPr is None:
        rPr = OxmlElement("w:rPr")
        run._element.insert(0, rPr)
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in ("ascii", "hAnsi", "cs"):
        rFonts.set(qn(f"w:{attr}"), "Arial")
    return run


def _insert_body_paragraph(anchor, text):
    p = _blank_p_after(anchor, space_before=Pt(0), space_after=Pt(6))
    _add_arial(p, text, size_pt=10)
    return p


def _insert_image(anchor, png_path, width_in):
    p = _blank_p_after(anchor, alignment=WD_ALIGN_PARAGRAPH.CENTER,
                       space_before=Pt(6), space_after=Pt(0),
                       line_spacing=1.15)
    run = p.add_run()
    run.add_picture(str(png_path), width=Inches(width_in))
    return p


def _insert_caption(anchor, text, italic=True):
    p = _blank_p_after(anchor, alignment=WD_ALIGN_PARAGRAPH.CENTER,
                       space_before=Pt(0), space_after=Pt(6))
    _add_arial(p, text, size_pt=10, italic=italic)
    return p


def _insert_label_paragraph(anchor, label, body):
    """Insert body paragraph with a bold leading label like 'BPR Loss. …'."""
    p = _blank_p_after(anchor, space_before=Pt(0), space_after=Pt(6))
    _add_arial(p, label + " ", size_pt=10, bold=True)
    _add_arial(p, body, size_pt=10)
    return p


# ---- content --------------------------------------------------------

ARCHITECTURE_PARAS = [
    (
        "Two embedding spaces, one shared user identity.",
        "CMF keeps three embedding matrices: a user matrix U of shape "
        "n_users × 96, a movie item matrix V_m of shape n_movies × 96, "
        "and a game item matrix V_g of shape n_games × 96. Unlike "
        "MF-BPR, which would train a completely separate U_movie and "
        "U_game, CMF uses the *same* U for both domains. Each user "
        "therefore has exactly one 96-dimensional vector, and that vector "
        "participates in both movie and game scoring."
    ),
    (
        "Item matrices stay private.",
        "Movies and games live in disjoint catalogues, so there is no "
        "reason for them to share an item matrix. V_m is updated only "
        "from movie interactions; V_g is updated only from game "
        "interactions. The shared U is what couples the two domains — "
        "every time a gradient flows through U from a movie triplet, it "
        "also shifts how that user scores games, and vice versa."
    ),
    (
        "Why 96 dimensions.",
        "The embedding width is the single architectural knob in CMF. "
        "Too small and the shared U cannot represent both tastes; too "
        "large and the two domain signals start to live in separate "
        "sub-spaces of U and the transfer benefit vanishes. We swept "
        "{32, 64, 96, 128} and picked 96 — large enough to hold two "
        "tastes without leaving enough idle capacity for them to "
        "decouple."
    ),
]


TRAINING_INTRO_PARAS = [
    (
        "Sampling.",
        "Each training step draws one BPR triplet per domain: "
        "(u, i_m^+, i_m^-) for movies and (u, i_g^+, i_g^-) for games, "
        "with negatives sampled uniformly at random. Only users who "
        "exist in the cohort are sampled; users with history in only "
        "one domain still contribute — their U[u] row is updated from "
        "that domain's triplets, just with no counter-pull from the "
        "other domain."
    ),
    (
        "Joint objective.",
        "The loss is a weighted sum of two per-domain BPR terms sharing "
        "U (Equation 4). α ∈ [0, 1] trades movie signal for game signal: "
        "at α = 0.5 the two domains contribute equal gradient; moving α "
        "toward 1 prioritises movie fit, toward 0 prioritises game fit. "
        "Our default is α = 0.5, and Lesson 4 sweeps α ∈ "
        "{0.3, 0.5, 0.7} to show the trade-off in isolation."
    ),
]


TRAINING_CLOSE_PARAS = [
    (
        "Optimisation.",
        "We initialise U, V_m, V_g with Gaussian noise (σ = 0.01), use "
        "SGD with learning rate 0.05, L2 regularisation λ = 1e-5 on Θ, "
        "and train for 30 epochs. Figure 13 visualises one step: the "
        "same U[u] row is pulled in two directions at once — toward "
        "movies this user liked and toward games this user liked — so "
        "the converged U is a compromise that must score well in both "
        "domains simultaneously."
    ),
]


INFERENCE_EXTRA_BODY = (
    "Figure 13b walks through the scoring path for a single user: their "
    "row U[u] is pulled out of the shared matrix, dot-producted against "
    "every row of V_game, and the resulting vector of scores is sorted "
    "to produce the top-K list. The movie item matrix V_m is not "
    "touched at inference — it has done its job by shaping U during "
    "training."
)


# ---- main -----------------------------------------------------------

def main():
    doc = Document(str(DOC))

    h_arch = _find_heading(doc, "CMF Architecture")
    h_train = _find_heading(doc, "CMF Training")
    h_infer = _find_heading(doc, "CMF Inference")
    h_limits = _find_heading(doc, "CMF Limitations")

    # --- Wipe Architecture body (between h_arch and h_train).
    _delete_between(doc, h_arch, h_train)
    # --- Wipe Training body (between h_train and h_infer).
    _delete_between(doc, h_train, h_infer)

    # --- Rebuild CMF Architecture (insert in reverse so order is preserved).
    for label, body in reversed(ARCHITECTURE_PARAS):
        # We rebuild by always inserting just after the Architecture heading.
        _insert_label_paragraph(h_arch, label, body)

    # --- Find the newly-inserted CMF Training heading (same element, but
    # neighbours changed) — we actually still have the h_train reference.

    # Insert a blank spacer paragraph right before CMF Training H4.
    # Not strictly needed — heading spacing handles it.

    # --- Rebuild CMF Training body below h_train, ending with Figure 13.
    # Build bottom-up (all anchored to h_train.addnext), so final top-down
    # order is: training intros → loss formula → where-explanation →
    # training close paragraphs → Figure 13 image + caption.

    # 6. Figure 13 caption (bottom).
    _insert_caption(
        h_train,
        "Figure 13: CMF mechanism — one shared user vector is pulled "
        "simultaneously by movie and game BPR triplets."
    )
    # 5. Figure 13 image.
    _insert_image(h_train, FIGS / "fig13_cmf.png", width_in=IMAGE_WIDTH_IN)

    # 4. Training close paragraphs.
    for label, body in reversed(TRAINING_CLOSE_PARAS):
        _insert_label_paragraph(h_train, label, body)

    # 3. Where-explanation paragraph for Equation 4.
    _insert_body_paragraph(
        h_train,
        "where U is the shared user embedding matrix used by both "
        "domains, V_m and V_g are the movie and game item matrices; "
        "L_BPR^movie and L_BPR^game are the per-domain BPR losses "
        "(same form as Equation 1) computed on each domain's triplets; "
        "α ∈ [0, 1] weights movie vs. game signal; Θ = {U, V_m, V_g} "
        "collects all trainable parameters; λ is the L2 regularisation "
        "strength. The shared U is the transfer mechanism — every "
        "gradient step on one domain shifts how that user scores items "
        "in the other."
    )

    # 2. Equation 4 caption.
    _insert_caption(
        h_train,
        "Equation 4: CMF joint two-domain loss with shared user matrix U."
    )
    # 1. Equation 4 image.
    _insert_image(h_train, FIGS / "fig_loss_cmf.png", width_in=FORMULA_WIDTH_IN)

    # 0. Training intro paragraphs (top of Training subsection).
    for label, body in reversed(TRAINING_INTRO_PARAS):
        _insert_label_paragraph(h_train, label, body)

    # --- CMF Inference: append new figure + caption AFTER the existing
    # body paragraph. We keep the existing body (para immediately after
    # h_infer). Insert after it — anchor on the existing body para.
    # Find the first non-empty paragraph after h_infer.
    from docx.text.paragraph import Paragraph
    cur_elem = h_infer._p.getnext()
    while cur_elem is not None and cur_elem.tag == f"{W}p":
        p = Paragraph(cur_elem, h_infer._parent)
        if p.text.strip():
            break
        cur_elem = cur_elem.getnext()
    if cur_elem is None or cur_elem.tag != f"{W}p":
        raise ValueError("no body paragraph found after CMF Inference heading")
    infer_body = Paragraph(cur_elem, h_infer._parent)

    # Insert caption, then image, then extra body — in reverse so top-down
    # order is: [existing body] → [extra body] → [image] → [caption].
    _insert_caption(
        infer_body,
        "Figure 13b: CMF inference — the shared user vector scores every "
        "game by dot product against V_game."
    )
    _insert_image(infer_body, FIGS / "fig13b_cmf_inference.png",
                  width_in=IMAGE_WIDTH_IN)
    _insert_body_paragraph(infer_body, INFERENCE_EXTRA_BODY)

    doc.save(str(DOC))
    print("CMF section rewritten.")


if __name__ == "__main__":
    main()
