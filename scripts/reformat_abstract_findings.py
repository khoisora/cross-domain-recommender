"""Replace the 'Headline findings' blob at p57 with a 'Key findings' list
where each numbered point sits on its own paragraph.
"""

from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parent.parent
DOC = ROOT / "project_report_v2.docx"


FINDINGS = [
    "Key findings.",
    "(1) User overlap is the decisive variable: on the 100%-overlap cohort (L3) "
    "LightGCN reaches Recall@10 = 0.059 and CDR models recover a +41% (EMCDR) to "
    "+750% (CMF/BiTGCF) lift over the low-overlap baseline (L2).",
    "(2) At cold-start (L6), single-domain models collapse (LightGCN 0.010), CDR "
    "models recover (EMCDR 0.033), and a simple popularity baseline tops the board "
    "at 0.038 — best-in-class.",
    "(3) A training-free co-occurrence reranking layer (λ = 0.05) improves every "
    "model on LLO (+31% on EMCDR) and rescues single-domain models at cold-start "
    "(LightGCN 0.010 → 0.033).",
    "(4) SBERT-CDR dominates niche items (bottom-50% popularity) by roughly 9× "
    "over collaborative baselines.",
    "The primary practical output is a context-aware routing rule — popularity + "
    "SBERT-CDR at zero game history, EMCDR/PTUPCDR at 1–2 games, LightGCN beyond "
    "— shipped as a nine-row hybrid demo.",
]


def set_paragraph_text(p, text):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    p.add_run(text)


def insert_paragraph_after(paragraph, text=None, style=None):
    new_p = paragraph._parent.add_paragraph(text or "", style=style)
    paragraph._p.addnext(new_p._p)
    return new_p


def main():
    doc = Document(DOC)
    ps = doc.paragraphs
    assert ps[57].text.startswith("Headline findings"), ps[57].text[:80]

    # Replace p57 with the first line ("Key findings.") then append the rest
    set_paragraph_text(ps[57], FINDINGS[0])
    # Insert subsequent lines in reverse so each lands right after p57 in order
    for line in reversed(FINDINGS[1:]):
        insert_paragraph_after(ps[57], line)

    doc.save(DOC)
    print(f"Saved {DOC}")


if __name__ == "__main__":
    main()
