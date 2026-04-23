"""Insert tables and remaining structural edits (P0 table + P1 items).

Adds:
  - Unified results table at end of §5
  - Dataset variant table in §3.2
  - Hyperparameter defaults table in §4.1
  - SBERT vs SBERT-CDR clarifying paragraph in §4.8
  - Related Work subsection at end of §4
  - Compact §7.6 overview table
"""

from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches

ROOT = Path(__file__).resolve().parent.parent
DOC = ROOT / "project_report_v2.docx"


def insert_paragraph_after(paragraph, text=None, style=None):
    new_p = paragraph._parent.add_paragraph(text or "", style=style)
    paragraph._p.addnext(new_p._p)
    return new_p


def insert_table_after(doc, anchor_para, data, style="Light Grid Accent 1"):
    """Create a table at end of body then move it to after anchor paragraph."""
    table = doc.add_table(rows=len(data), cols=len(data[0]))
    try:
        table.style = style
    except KeyError:
        pass
    for i, row in enumerate(data):
        for j, cell in enumerate(row):
            c = table.rows[i].cells[j]
            c.text = str(cell)
            # Bold header row
            if i == 0:
                for p in c.paragraphs:
                    for r in p.runs:
                        r.bold = True
            # Shrink font in all cells
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    anchor_para._p.addnext(table._tbl)
    return table


# ─── table content ─────────────────────────────────────────────────────

UNIFIED_RESULTS = [
    ["Model", "L1", "L2", "L3", "L3+cooc", "L4", "L5", "L6", "L6+cooc", "L7"],
    ["Popularity",     "—",     "—",     "—",     "—",     "—",     "—",     "0.038", "0.039", "—"],
    ["MF-BPR",         "0.032", "0.017", "0.045", "0.052", "0.010", "0.004", "0.001", "0.020", "—"],
    ["NCF",            "0.016", "0.012", "0.026", "0.037", "0.023", "0.018", "0.002", "—",     "—"],
    ["LightGCN",       "0.032", "0.029", "0.060", "0.063", "0.038", "0.037", "0.010", "0.033", "0.034"],
    ["CMF",            "0.031", "0.005", "0.040", "0.038", "0.014", "0.016", "0.001", "0.032", "—"],
    ["EMCDR",          "0.021", "0.016", "0.023", "0.034", "0.027", "0.010", "0.033", "0.034", "—"],
    ["PTUPCDR",        "0.025", "0.009", "0.032", "0.036", "0.032", "0.010", "0.030", "0.030", "0.022"],
    ["BiTGCF",         "—",     "0.005", "0.043", "0.053", "0.023", "0.028", "0.004", "0.025", "—"],
    ["SBERT",          "0.031", "—",     "—",     "—",     "—",     "—",     "—",     "—",     "0.031"],
    ["SBERT-CDR",      "—",     "—",     "—",     "—",     "—",     "—",     "0.002", "0.024", "0.033"],
]

DATASET_VARIANTS = [
    ["Cohort", "Lessons", "Definition", "Size", "Isolates"],
    ["Base", "L2, L8-LLO",
     "movie_game k-core ≥ 10",
     "≈14.3k users\n≈389k movie / ≈59k game interactions",
     "baseline CDR in the wild"],
    ["100% overlap", "L3",
     "users with ≥ 1 rating in both domains",
     "≈19.9k users",
     "user-overlap ratio"],
    ["Source-rich / target-sparse", "L4",
     "overlap users with ≥ 10 movie ratings and ≤ 5 game ratings",
     "≈5.0k users",
     "source-density effect"],
    ["Catalog-sharpened", "L5",
     "base cohort with long-tail items (< 5 interactions) pruned",
     "≈14k users, smaller item catalog",
     "noise from long-tail items"],
    ["Cold-start split", "L6, L8-cold",
     "80/20 user split; held-out 20% have zero game ratings at train time",
     "≈11.5k train / ≈2.9k cold-test users",
     "true cold-start regime"],
    ["Niche-item subgroups", "L7",
     "base cohort partitioned by target-item popularity",
     "bottom-50% vs top-50% popularity splits",
     "content vs collaborative on niche items"],
]

HYPERPARAMS = [
    ["Model", "Key hyperparameters", "Default (this report)", "Tuned on", "Sensitivity"],
    ["MF-BPR",   "embedding_dim, lr, neg-samples",        "64 / 1e-3 / 4",              "—",   "low"],
    ["NCF",      "GMF dim, MLP layers, lr, dropout",       "64 / [128,64,32,16] / 1e-3 / 0.2", "—", "low"],
    ["LightGCN", "K (propagation depth), embedding_dim",   "K = 4, dim = 64",             "L3 (Fig. 37)", "medium — plateau after K ≥ 2"],
    ["CMF",      "lr, α (source-loss weight)",             "lr = 1e-3, α = 0.2",          "L3 (Fig. 35)", "high — default α = 0.5 halved Recall@10"],
    ["EMCDR",    "mapping-MLP layers, λ (cooc blend)",     "[128,64], λ = 0.05",          "L3 (Fig. 36)", "high on λ — peaks sharply at 0.05"],
    ["PTUPCDR",  "n_experts, gate τ, blend w",             "n = 2, τ = 1.0, w = 0.5",     "L3 (Fig. 38)", "medium — more experts over-fit L3"],
    ["SBERT-CDR","source_weight (movie share)",            "0.1 on overlap, 0.5 on cold",  "L3 (Fig. 39)", "high — monotonic on overlap cohorts"],
    ["Cooc rerank","λ (blend with base score)",            "0.05",                        "L3 (§5.8)",    "high past λ = 0.1 — personalisation collapses"],
]

SECTION_7_6_ROWS = [
    ["#", "Row", "Model", "Lesson motivating it"],
    ["1", "Top Picks (primary)", "LightGCN / EMCDR / Popularity (routed)", "L6 cold-start + L3 overlap"],
    ["2", "Also Try (complementary)", "PTUPCDR or SBERT-CDR", "L4 sparse-target / L7 niche"],
    ["3", "Movie-Fans Also Played (cooc)", "Co-occurrence rerank on popularity", "L8 universal rerank"],
    ["4", "Hidden Gems", "SBERT-CDR on bottom-50% popularity items", "L7 niche subgroup"],
    ["5", "Because You Liked (movies)", "SBERT item-item", "L1 content baseline"],
    ["6", "More Like Your Movies", "LightGCN-Movies", "L2 in-domain backbone"],
    ["7", "Games Movie-Fans Play", "Reverse co-occurrence", "L8 symmetric transfer"],
    ["8", "Popular Games", "Popularity", "L6 popularity ceiling"],
    ["9", "Popular Movies", "Popularity", "L6 same-domain safety baseline"],
]


# ─── prose edits ───────────────────────────────────────────────────────

SBERT_CLARIFIER = (
    "A note on naming. We distinguish between SBERT (content-only, single-domain: "
    "movie-to-movie similarity in the semantic space, used as the in-domain baseline "
    "in Lessons 1 and 7) and SBERT-CDR (the cross-domain variant described here: a "
    "user profile is formed from their rated items — movies and, where available, "
    "games — and scored against game item embeddings in the same shared 384-dim "
    "space). Only SBERT-CDR crosses domains; plain SBERT is included as a reference "
    "to show that collaborative models beat content similarity when target-domain "
    "history is available."
)

RELATED_WORK_PARAGRAPHS = [
    ("4.10 Related Work", "Heading 2"),
    (
        "Cross-domain recommendation has been surveyed comprehensively by Zhu et al. "
        "[9] and shares roots with transfer learning. Three lines of prior work are "
        "most directly relevant. First, mapping-based CDR: EMCDR [4] and PTUPCDR [5] "
        "train a function that maps source-domain user embeddings into the target "
        "domain — we port both as-is. Second, joint-training CDR: CMF [3], CoNet "
        "[11], and DDTCDR [10] learn shared representations across domains; we port "
        "CMF as the simplest member of this family and include BiTGCF [12] as a "
        "graph-based analogue. Third, content-based bridging: LLM-derived text "
        "embeddings — Sentence-BERT [7] and dataset-specific variants such as "
        "BLaIR [8] — enable domain-agnostic semantic matching. Our SBERT-CDR is an "
        "off-the-shelf instantiation of this idea. Orthogonal to model choice, "
        "graph-convolution baselines [2, 14], BPR [6], and factorisation machines "
        "[15] bound the single-domain ranker; our LightGCN and MF-BPR numbers are "
        "reproductions under the same leave-last-out protocol. We also note sampled-"
        "metric pitfalls [16] — our reported numbers are full-rank Recall@10 / "
        "NDCG@10 and therefore avoid the biases documented there.",
        None,
    ),
]


# ─── main ─────────────────────────────────────────────────────────────

def main():
    doc = Document(DOC)
    ps = doc.paragraphs

    # Re-locate anchors (post-first-pass indices).
    def find(needle, style_substr):
        for i, p in enumerate(ps):
            if needle in p.text and style_substr in p.style.name:
                return i
        raise LookupError(f"{needle!r} not found with style {style_substr!r}")

    idx_32 = find("3.2 Dataset Variants", "Heading 2")                # 100
    idx_33 = find("3.3 Data Splitting Protocol", "Heading 2")          # 108
    idx_41 = find("4.1 Evaluation Protocol", "Heading 2")              # 113
    idx_42 = find("4.2 MF-BPR", "Heading 2")                           # 119
    idx_48 = find("4.8 SBERT-CDR", "Heading 2")                        # 172
    idx_49 = find("4.9 Co-occurrence", "Heading 2")                    # 181
    idx_59 = find("5.9 Hyperparameter", "Heading 2")                   # 227
    idx_76 = find("7.6 Model Selection", "Heading 2")                  # 314

    # Bottom-up insertion order.

    # 1. §7.6 compact overview table — insert right after the §7.6 heading intro
    #    (the paragraph right after the heading). Target = idx_76 + 1.
    overview_anchor = ps[idx_76 + 1]  # "Section 7.2 described what each of the nine..."
    # Add a short caption before the table
    cap = insert_paragraph_after(overview_anchor,
                                  "Table 5: Nine-row demo overview — each row maps to the lesson that motivated it.",
                                  style="Caption")
    # Now insert the table AFTER the caption (so caption sits above the table)
    insert_table_after(doc, cap, SECTION_7_6_ROWS)

    # 2. Unified results table — insert BEFORE §5.9 Hyperparameter heading.
    #    Anchor is the paragraph immediately preceding idx_59.
    before_59 = ps[idx_59 - 1]
    # Insert caption then table just after this anchor (so they sit before §5.9).
    cap = insert_paragraph_after(before_59,
                                  "Table 4: Unified Recall@10 across all lessons and models "
                                  "(— = not evaluated in that lesson).",
                                  style="Caption")
    insert_table_after(doc, cap, UNIFIED_RESULTS)

    # 3. Related Work — after §4.9 block. Find the last paragraph of §4.9 before §5.
    idx_5 = find("5. Experiments", "Heading 1")
    anchor_after_49 = ps[idx_5 - 1]  # last paragraph of §4.9 content
    for text, style in reversed(RELATED_WORK_PARAGRAPHS):
        insert_paragraph_after(anchor_after_49, text, style=style)

    # 4. SBERT clarifier — insert right after the §4.8 heading.
    insert_paragraph_after(ps[idx_48], SBERT_CLARIFIER)

    # 5. Hyperparameter defaults table — insert right before §4.2.
    before_42 = ps[idx_42 - 1]
    cap = insert_paragraph_after(before_42,
                                  "Table 3: Key hyperparameters, defaults adopted in this report, "
                                  "and sensitivity (tuned values are from the §5.9 sweeps).",
                                  style="Caption")
    insert_table_after(doc, cap, HYPERPARAMS)

    # 6. Dataset variants table — insert right before §3.3.
    before_33 = ps[idx_33 - 1]
    cap = insert_paragraph_after(before_33,
                                  "Table 2: Dataset cohorts used across the lesson series.",
                                  style="Caption")
    insert_table_after(doc, cap, DATASET_VARIANTS)

    doc.save(DOC)
    print(f"Saved {DOC}")


if __name__ == "__main__":
    main()
