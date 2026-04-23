"""Apply uniform spacing to every Heading level.

After moving SBERT/Cooc blocks, Heading 3 and Heading 4 paragraphs have a
mix of space_before/space_after values — some inherited, some explicit.
This pass sets one canonical pair per level so spacing is consistent.

Line spacing stays at 1.35 everywhere.
"""

from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_LINE_SPACING

ROOT = Path(__file__).resolve().parent.parent
DOC = ROOT / "project_report_v2.docx"

# (space_before, space_after) in points per heading level.
HEADING_SPACING = {
    "Heading 1": (Pt(18), Pt(10)),
    "Heading 2": (Pt(12), Pt(8)),
    "Heading 3": (Pt(10), Pt(6)),
    "Heading 4": (Pt(6), Pt(6)),
}


def _apply(p):
    style = p.style.name
    if style not in HEADING_SPACING:
        return False
    sb, sa = HEADING_SPACING[style]
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.35
    pf.space_before = sb
    pf.space_after = sa
    return True


def main():
    doc = Document(str(DOC))
    count = 0
    for p in doc.paragraphs:
        if _apply(p):
            count += 1
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if _apply(p):
                        count += 1
    doc.save(str(DOC))
    print(f"Normalized spacing on {count} heading paragraphs")


if __name__ == "__main__":
    main()
