"""Compare §4 structure/styling between redesigned and current docs."""

from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
DOC_NEW = ROOT / "project_report_v2_redesigned.docx"
DOC_OLD = ROOT / "project_report_v2.docx"


def _dump_s4(doc, label):
    print(f"\n{'='*70}")
    print(f"  {label}")
    print(f"{'='*70}")
    in_s4 = False
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        s = p.style.name
        sid = p.style.style_id

        if s == "Heading 1":
            if t.startswith("4.") or t.startswith("4 "):
                in_s4 = True
            elif in_s4:
                break
        if not in_s4:
            continue

        # Run info
        first_r = p._p.find(qn("w:r"))
        run_info = ""
        if first_r is not None:
            rPr = first_r.find(qn("w:rPr"))
            if rPr is not None:
                font_el = rPr.find(qn("w:rFonts"))
                font = font_el.get(qn("w:ascii"), "") if font_el is not None else ""
                sz_el = rPr.find(qn("w:sz"))
                sz = sz_el.get(qn("w:val"), "") if sz_el is not None else ""
                b_el = rPr.find(qn("w:b"))
                bold = "B" if b_el is not None else ""
                i_el = rPr.find(qn("w:i"))
                ital = "I" if i_el is not None else ""
                color_el = rPr.find(qn("w:color"))
                color = f"#{color_el.get(qn('w:val'))}" if color_el is not None else ""
                run_info = f" [{font} {sz}{bold}{ital} {color}]"

        # Spacing
        pPr = p._p.find(qn("w:pPr"))
        sp_info = ""
        if pPr is not None:
            sp = pPr.find(qn("w:spacing"))
            if sp is not None:
                b = sp.get(qn("w:before"), "")
                a = sp.get(qn("w:after"), "")
                l = sp.get(qn("w:line"), "")
                sp_info = f" sp(b={b},a={a},l={l})"

        print(f"p#{i} [{s}/{sid}]{sp_info}{run_info}")
        print(f"     {t[:130]}")


def main():
    doc_new = Document(str(DOC_NEW))
    doc_old = Document(str(DOC_OLD))
    _dump_s4(doc_new, "REDESIGNED §4")
    _dump_s4(doc_old, "CURRENT §4")


if __name__ == "__main__":
    main()
