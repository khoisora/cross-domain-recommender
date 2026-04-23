"""Replace the embedded Figure L1 image in project_report_v2.docx."""

from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
DOC = ROOT / "project_report_v2.docx"
NEW_IMG = ROOT / "report_figures_v3" / "fig_lesson1_pointwise_pairwise.png"


def _para_text(p):
    return "".join(t.text or "" for t in p._p.iter(qn("w:t")))


def main():
    doc = Document(str(DOC))

    # Find the paragraph containing the Figure L1 image.
    target_p = None
    for i, p in enumerate(doc.paragraphs):
        drawings = p._p.findall(".//" + qn("w:drawing"))
        if drawings and i + 1 < len(doc.paragraphs):
            nxt = _para_text(doc.paragraphs[i + 1]).strip()
            if "Figure L1" in nxt:
                target_p = p
                print(f"Found Figure L1 image at p#{i}")
                break

    if target_p is None:
        print("Figure L1 not found!")
        return

    # Find the blip and its rId.
    ns_a = "http://schemas.openxmlformats.org/drawingml/2006/main"
    ns_r = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    blip = target_p._p.find(f".//{{{ns_a}}}blip")
    old_rid = blip.get(f"{{{ns_r}}}embed")
    print(f"  Old rId: {old_rid}")

    # Get the image part referenced by this rId and replace its blob.
    part = doc.part
    rel = part.rels[old_rid]
    image_part = rel.target_part
    print(f"  Old image content_type: {image_part.content_type}")

    # Replace the blob with the new image bytes.
    new_bytes = NEW_IMG.read_bytes()
    image_part._blob = new_bytes
    print(f"  Replaced with {len(new_bytes)} bytes from {NEW_IMG.name}")

    doc.save(str(DOC))
    print("✓ Figure L1 image replaced in docx")


if __name__ == "__main__":
    main()
