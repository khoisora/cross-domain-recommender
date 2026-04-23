"""Replace every Experiment Setup / Key Finding callout table with a plain
body paragraph that carries a bold inline label.

Before (rendered as a yellow 1×1 or 2×1 table):

    ┌──────────────────────────────────────────────┐
    │ Experiment Setup                             │
    │ Variable Changed: … Held Constant: …         │
    └──────────────────────────────────────────────┘

After (plain body paragraph, matches the "Question. …" style already
used elsewhere in the report):

    **Experiment Setup.** Variable Changed: … Held Constant: …

For tables that also have a Key Finding row, that row is emitted as a
second paragraph with the same treatment.

The body paragraph's XML is cloned wholesale so internal line breaks
(w:br) and run styling are preserved — we only prepend a new bold run
carrying the label.
"""

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
DOC = ROOT / "project_report_v2.docx"

LABELS = ("Experiment Setup", "Key Finding")


def _cell_text(cell):
    return " ".join(p.text for p in cell.paragraphs).strip()


def _make_bold_label_run(text):
    """Build a <w:r> carrying the bold label (Arial 10pt, bold)."""
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")

    rFonts = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "cs"):
        rFonts.set(qn(f"w:{attr}"), "Arial")
    rPr.append(rFonts)

    b = OxmlElement("w:b")
    rPr.append(b)

    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "20")  # 10pt
    rPr.append(sz)
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), "20")
    rPr.append(szCs)

    r.append(rPr)

    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    return r


def _build_paragraph_from_cell(label, body_para):
    """Clone body_para's <w:p>, prepend a bold label run, return new element."""
    new_p = deepcopy(body_para._p)
    # Find first w:r child (the actual run content); prepend the label run
    # in front of all existing runs so it renders as "Label. <body>".
    first_r = new_p.find(qn("w:r"))
    label_run = _make_bold_label_run(f"{label}. ")
    if first_r is not None:
        first_r.addprevious(label_run)
    else:
        new_p.append(label_run)
    return new_p


def main():
    doc = Document(str(DOC))

    # Identify callout tables: the first cell's first paragraph text must
    # equal one of our labels.
    targets = []
    for t in doc.tables:
        try:
            first_para_text = t.rows[0].cells[0].paragraphs[0].text.strip()
        except IndexError:
            continue
        if first_para_text in LABELS:
            targets.append(t)

    print(f"Found {len(targets)} Experiment-Setup-style callout tables")

    unboxed = 0
    for t in targets:
        tbl = t._tbl
        parent = tbl.getparent()
        insert_idx = list(parent).index(tbl)

        for row in t.rows:
            cell = row.cells[0]
            if len(cell.paragraphs) < 2:
                continue
            label = cell.paragraphs[0].text.strip().rstrip(":. ")
            body_p = cell.paragraphs[1]
            new_p = _build_paragraph_from_cell(label, body_p)
            parent.insert(insert_idx, new_p)
            insert_idx += 1

        parent.remove(tbl)
        unboxed += 1
        print(f"  ✓ unboxed: {_cell_text(t.rows[0].cells[0])[:80]}…")

    doc.save(str(DOC))
    print(f"\nUnboxed {unboxed} callout tables")


if __name__ == "__main__":
    main()
