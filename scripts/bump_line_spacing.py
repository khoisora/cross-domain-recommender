"""Increase line spacing for readability.

Sets line_spacing = 1.35 (multiple) on every body paragraph, including those
inside tables. Existing space_before/space_after are left unchanged.
"""

from pathlib import Path
from docx import Document
from docx.enum.text import WD_LINE_SPACING

ROOT = Path(__file__).resolve().parent.parent
DOC = ROOT / "project_report_v2.docx"

TARGET_SPACING = 1.35


def _apply(p):
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = TARGET_SPACING


def main():
    doc = Document(str(DOC))
    count = 0
    for p in doc.paragraphs:
        _apply(p)
        count += 1
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _apply(p)
                    count += 1
    doc.save(str(DOC))
    print(f"Updated line spacing on {count} paragraphs → {TARGET_SPACING}")


if __name__ == "__main__":
    main()
