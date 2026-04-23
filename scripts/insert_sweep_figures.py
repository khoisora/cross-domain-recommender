"""Insert §5.9 hyperparameter-sweep figures and refresh body text.

For each of CMF / EMCDR-cooc / LightGCN / PTUPCDR / SBERT-CDR:
  - replace the body paragraph's runs with the updated description
  - insert an image paragraph + caption after it
"""

from pathlib import Path
from docx import Document
from docx.shared import Inches
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
DOC = ROOT / "project_report_v2.docx"
FIG_DIR = ROOT / "report_figures_v3"


def set_paragraph_text(p, text):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    p.add_run(text)


def insert_paragraph_after(paragraph, text=None, style=None):
    new_p = paragraph._parent.add_paragraph(text or "", style=style)
    paragraph._p.addnext(new_p._p)
    return new_p


def insert_figure_after(paragraph, image_path, caption, width_inches=6.0):
    img_p = insert_paragraph_after(paragraph)
    img_p.alignment = 1  # CENTER
    run = img_p.add_run()
    run.add_picture(str(image_path), width=Inches(width_inches))
    cap_p = insert_paragraph_after(img_p, caption, style="Caption")
    cap_p.alignment = 1
    return cap_p


# ─── new body text aligned to actual sweep results ──────────────────────

NEW_CMF = (
    "With the defaults inherited from the reference implementation, CMF reached only "
    "Recall@10 ≈ 0.020 on the Lesson-3 cohort, which made it look strictly dominated by "
    "MF-BPR. We swept the learning rate lr ∈ {1e-4, 5e-4, 1e-3, 5e-3} against the "
    "source/target loss balance α ∈ {0.01, 0.05, 0.2, 0.5} — 16 configurations in total "
    "(Figure 22). The heatmap reveals a single sweet spot at lr = 1e-3, α = 0.2 "
    "(Recall@10 = 0.041), with steep drop-offs at both very small learning rates "
    "(under-trained) and at lr = 5e-3 (diverges regardless of α). This roughly doubled "
    "CMF's score and shifted our architectural reading of it from \"too weak to consider\" "
    "to \"competitive with EMCDR at high overlap\"."
)

NEW_EMCDR = (
    "The base EMCDR numbers understated the model's ceiling in our early runs. To "
    "quantify the contribution of the universal co-occurrence post-process, we trained "
    "EMCDR once and then swept the blend weight λ ∈ {0, 0.02, 0.05, 0.1, 0.2, 0.3} "
    "(Figure 23). Both Recall@10 and NDCG@10 peak sharply at λ = 0.05 (Recall@10 = 0.034 "
    "vs 0.026 at λ = 0, a +31% lift); past λ = 0.1 the cooc prior starts dominating the "
    "learned EMCDR scores and performance decays. Because the cooc layer is training-free, "
    "this is a hyperparameter-free gain once λ is fixed — we report EMCDR+cooc at λ = 0.05 "
    "as the default serving configuration throughout Section 7."
)

NEW_LIGHTGCN = (
    "We swept K ∈ {1, 2, 3, 4, 5} propagation layers on the Lesson-3 cohort (Figure 24). "
    "Recall@10 climbs sharply from K = 1 (0.045) to K = 2 (0.056), then plateaus: "
    "K = 3 and K = 4 differ by less than 0.001, with K = 4 nominally best at "
    "Recall@10 = 0.057. NDCG@10 peaks jointly at K = 2 and K = 4 (0.030). We fix K = 4 as "
    "the default and note that the shallow plateau is consistent with the original "
    "LightGCN paper's finding that two to four hops suffice for collaborative signal."
)

NEW_PTUPCDR = (
    "We swept n_experts ∈ {2, 4, 8, 16} per-user MLPs in the meta-network on the "
    "Lesson-3 cohort (Figure 25). Counter-intuitively, fewer experts won: "
    "n_experts = 2 achieves Recall@10 = 0.033 and NDCG@10 = 0.018, versus 8 experts at "
    "0.030 / 0.015 and 16 experts at 0.029 / 0.015. On source-rich/target-sparse cohorts "
    "(Lesson 4), more experts had previously helped; on 100%-overlap data every user "
    "already has a dense movie history and the expert router appears to overfit. The "
    "gate softmax temperature was left at the reference default (τ = 1.0); sweeping it "
    "(0.5–2.0) did not produce a reliable improvement."
)

NEW_SBERT = (
    "The SBERT-CDR user profile is a weighted mean of rated-movie and rated-game item "
    "embeddings; source_weight sets the movie share of that blend. We swept "
    "source_weight ∈ {0.1, 0.3, 0.5, 0.7, 0.9} (Figure 26). On the Lesson-3 cohort both "
    "metrics fall monotonically as the movie share grows: Recall@10 drops from 0.029 at "
    "0.1 to 0.004 at 0.9, with a sharp cliff between 0.5 and 0.7. Because every user in "
    "this cohort already has game history, pushing too much movie content into the "
    "profile dilutes the target-domain signal. We fix source_weight = 0.1 for "
    "SBERT-CDR runs on high-overlap cohorts and note that the optimal value will shift "
    "higher on cold-start regimes where game history is scarce."
)


def main():
    doc = Document(DOC)
    ps = doc.paragraphs

    # Sanity-check the heading/body layout we located earlier.
    assert "CMF learning rate" in ps[227].text, ps[227].text
    assert "LightGCN propagation" in ps[231].text, ps[231].text

    # 1. Update body text in-place.
    set_paragraph_text(ps[228], NEW_CMF)
    set_paragraph_text(ps[230], NEW_EMCDR)
    set_paragraph_text(ps[232], NEW_LIGHTGCN)
    set_paragraph_text(ps[234], NEW_PTUPCDR)
    set_paragraph_text(ps[236], NEW_SBERT)

    # 2. Insert figures AFTER the body paragraphs — go bottom-up so earlier
    #    indices stay valid.
    insertions = [
        (236, FIG_DIR / "fig_hp_sbert_cdr.png",
         "Figure 26: SBERT-CDR source_weight sweep (Lesson 3 cohort) — best at source_weight = 0.1."),
        (234, FIG_DIR / "fig_hp_ptupcdr.png",
         "Figure 25: PTUPCDR n_experts sweep (Lesson 3 cohort) — best at n_experts = 2."),
        (232, FIG_DIR / "fig_hp_lightgcn.png",
         "Figure 24: LightGCN propagation-depth sweep (Lesson 3 cohort) — Recall@10 plateaus at K ≥ 2, best at K = 4."),
        (230, FIG_DIR / "fig_hp_emcdr_cooc.png",
         "Figure 23: EMCDR + co-occurrence rerank λ sweep — Recall@10 peaks at λ = 0.05 with a +31% lift over λ = 0."),
        (228, FIG_DIR / "fig_hp_cmf.png",
         "Figure 22: CMF lr × α grid (16 configurations on the Lesson-3 cohort) — Recall@10 peaks at lr = 1e-3, α = 0.2."),
    ]
    for idx, img, caption in insertions:
        insert_figure_after(ps[idx], img, caption)

    out = DOC
    doc.save(out)
    print(f"Saved {out}")


if __name__ == "__main__":
    main()
