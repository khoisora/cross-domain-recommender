"""Normalize paragraph spacing so body rhythm is uniform.

Line spacing is already 1.35 everywhere (from `bump_line_spacing.py`). The
visible inconsistency comes from `space_before` / `space_after` varying per
paragraph (0, 2.5pt, 4pt, 6pt, 12pt, 18pt values all present on `normal`).

This pass sets every `normal` paragraph to space_before = 0 and space_after
= 6pt, including paragraphs inside table cells. Heading styles keep their
own spacing.
"""

from pathlib import Path
from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent.parent
DOC = ROOT / "project_report_v2.docx"

BODY_SPACE_BEFORE = Pt(0)
BODY_SPACE_AFTER = Pt(6)


def _apply(p):
    if p.style.name != "normal":
        return False
    pf = p.paragraph_format
    pf.space_before = BODY_SPACE_BEFORE
    pf.space_after = BODY_SPACE_AFTER
    return True


def main():
    doc = Document(str(DOC))
    count = 0
    for p in doc.paragraphs:
        if _apply(p):
            count += 1
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if _apply(p):
                        count += 1
    doc.save(str(DOC))
    print(f"Normalized spacing on {count} normal paragraphs → {DOC}")


if __name__ == "__main__":
    main()
