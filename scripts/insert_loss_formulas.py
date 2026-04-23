"""Insert one rendered loss-formula PNG + caption + symbol explanation after
the paragraph that first introduces each model's loss.

For every entry in LOSS_SPECS we:
  1. Locate the anchor paragraph by text-prefix (unique).
  2. Insert, immediately after it, three new paragraphs:
       a. centered image paragraph (inline PNG at IMAGE_WIDTH_IN)
       b. italic, centered caption paragraph ("Equation N: ...")
       c. a "where" paragraph spelling out every symbol
  3. Existing surrounding prose is preserved — the formula + explanation
     slot naturally in front of the behavioural discussion already in the
     body.
"""

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
DOC = ROOT / "project_report_v2.docx"
FIG_DIR = ROOT / "report_figures_v3"

IMAGE_WIDTH_IN = 5.8


# Ordered so equation numbers follow Section-4 reading order,
# then §5.7 SBERT, then §5.8 Co-occurrence.
LOSS_SPECS = [
    dict(
        eq_num=1,
        anchor="BPR Loss. The Bayesian Personalized Ranking",
        png="fig_loss_mfbpr.png",
        caption="Equation 1: MF-BPR pairwise ranking loss.",
        explanation=(
            "where D is the set of training triplets (u, i⁺, i⁻) — one observed "
            "positive i⁺ paired with a uniformly sampled negative i⁻ for user u; "
            "u_u and v_i are the learned 64-dim user and item embeddings; "
            "σ(·) is the logistic sigmoid; Θ = {U, V} denotes all trainable "
            "parameters; and λ is the L2 regularisation strength. Minimising "
            "this loss pushes the positive-item score above the negative-item "
            "score by as wide a margin as the data supports."
        ),
    ),
    dict(
        eq_num=2,
        anchor="NCF trains using Binary Cross-Entropy",
        png="fig_loss_ncf.png",
        caption="Equation 2: NCF binary cross-entropy loss on observed positives and sampled negatives.",
        explanation=(
            "where D⁺ is the set of observed interactions (y_{u,i}=1) and "
            "D⁻ is a matched set of uniformly sampled negatives (y_{u,i}=0); "
            "ŷ_{u,i} ∈ (0, 1) is the model's predicted interaction probability, "
            "produced by a sigmoid on top of the fused GMF + MLP representation. "
            "BCE penalises confident wrong predictions sharply: a predicted 0.9 "
            "on a true negative contributes −log(0.1) ≈ 2.3 to the loss."
        ),
    ),
    dict(
        eq_num=3,
        anchor="LightGCN uses the same BPR loss as MF-BPR",
        png="fig_loss_lightgcn.png",
        caption="Equation 3: LightGCN layer-combined embeddings and BPR loss.",
        explanation=(
            "where e_v^{(k)} is node v's embedding after k graph-propagation "
            "steps and ẽ_v is the final representation — the unweighted mean "
            "across K+1 layers (we use K = 3). The propagation itself has no "
            "learnable weights; only the layer-0 embeddings E^{(0)} are "
            "trained, and the same BPR pairwise log-sigmoid objective is "
            "applied to the propagated user/item vectors."
        ),
    ),
    dict(
        eq_num=4,
        anchor="CMF maintains three embedding matrices",
        png="fig_loss_cmf.png",
        caption="Equation 4: CMF joint two-domain loss with shared user matrix U.",
        explanation=(
            "where U is the shared user embedding matrix used by both domains, "
            "V_m and V_g are the movie and game item matrices; "
            "L_BPR^{movie} and L_BPR^{game} are BPR losses computed on each "
            "domain's interactions; α ∈ [0, 1] weights the movie domain "
            "against the game domain (we sweep α in Lesson 4). The shared U "
            "is the transfer mechanism — optimising it on movie triplets "
            "moves the same user vectors that score game items."
        ),
    ),
    dict(
        eq_num=5,
        anchor="EMCDR trains sequentially in three distinct phases",
        png="fig_loss_emcdr.png",
        caption="Equation 5: EMCDR phase-3 mapping loss (MSE on overlap users).",
        explanation=(
            "where U_overlap is the set of users with interactions in both "
            "domains; u_u^{movie} and u_u^{game} are the user embeddings "
            "learned independently in phase 1 and phase 2; f_θ is the MLP "
            "mapping network trained in phase 3 to regress the game-side "
            "vector from the movie-side vector. At inference, a cold-start "
            "user's movie embedding is pushed through f_θ to get a usable "
            "game-space vector."
        ),
    ),
    dict(
        eq_num=6,
        anchor="Like EMCDR, PTUPCDR starts with Phase 1",
        png="fig_loss_ptupcdr.png",
        caption="Equation 6: PTUPCDR mixture-of-experts personalised mapping loss.",
        explanation=(
            "where each f_k is one of K expert MLPs (K = 8 in our runs) and "
            "g_k is the corresponding gate — a softmax-normalised scalar "
            "derived from the user's own movie embedding. The mapping is "
            "therefore personalised: the same movie-side vector is routed "
            "through different expert combinations for different users, "
            "which captures user-specific cross-domain preferences that a "
            "single global f_θ (as in EMCDR) cannot."
        ),
    ),
    dict(
        eq_num=7,
        anchor="The key insight is that movies and games are encoded",
        png="fig_loss_sbert.png",
        caption="Equation 7: SBERT-CDR user profile and cosine scoring (no training).",
        explanation=(
            "where I_u is the set of items user u has rated, e_j is the "
            "SBERT embedding (384-dim) of item j produced by the frozen "
            "Sentence-BERT encoder, and p_u is the user profile — the "
            "unweighted mean of their rated-item embeddings. Game items "
            "are scored by cosine similarity s_{u,i} against the profile. "
            "Nothing is trained in this model; all signal comes from the "
            "shared semantic space of movie and game descriptions."
        ),
    ),
    dict(
        eq_num=8,
        anchor="For every (movie, game) pair in the training data",
        png="fig_loss_cooc.png",
        caption="Equation 8: Co-occurrence matrix and reranking score.",
        explanation=(
            "where U_m and U_g are the sets of users who interacted with "
            "movie m and game g; c_{m,g} is their cross-domain co-occurrence, "
            "normalised by the geometric mean of the two audience sizes to "
            "prevent popular items from dominating. The reranked score "
            "s^{rerank}_{u,g} adds a weighted sum of c_{m,g} over the user's "
            "movie history to whatever base-model score s^{base}_{u,g} we "
            "started with; λ controls the rerank strength (we use λ = 0.05)."
        ),
    ),
]


def _find_paragraph(doc, prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    return None


def _make_blank_para_after(anchor_para):
    """Clone anchor's paragraph-properties element (pPr) into a brand-new
    empty paragraph and insert it immediately after anchor_para. Returns
    the python-docx Paragraph wrapper for the new element."""
    from docx.text.paragraph import Paragraph

    new_p = OxmlElement("w:p")
    # Copy pPr from anchor so style/spacing are compatible.
    src_pPr = anchor_para._p.find(qn("w:pPr"))
    if src_pPr is not None:
        new_p.append(deepcopy(src_pPr))
    anchor_para._p.addnext(new_p)
    return Paragraph(new_p, anchor_para._parent)


def _set_arial(run, size_pt=10, italic=False):
    run.font.name = "Arial"
    run.font.size = Pt(size_pt)
    run.italic = italic
    # Also set the complex-script / hAnsi slots.
    rPr = run._element.find(qn("w:rPr"))
    if rPr is not None:
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.insert(0, rFonts)
        for attr in ("ascii", "hAnsi", "cs"):
            rFonts.set(qn(f"w:{attr}"), "Arial")


def _insert_block(doc, anchor_para, png_path, caption_text, explanation_text):
    # Build paragraphs in reverse so each addnext keeps the anchor's
    # direct successor correct. We construct: anchor → [img] → [cap] → [exp].
    exp_p = _make_blank_para_after(anchor_para)
    exp_p.paragraph_format.line_spacing = 1.35
    exp_p.paragraph_format.space_before = Pt(0)
    exp_p.paragraph_format.space_after = Pt(6)
    run = exp_p.add_run(explanation_text)
    _set_arial(run, size_pt=10)

    cap_p = _make_blank_para_after(anchor_para)
    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_p.paragraph_format.line_spacing = 1.35
    cap_p.paragraph_format.space_before = Pt(0)
    cap_p.paragraph_format.space_after = Pt(6)
    run = cap_p.add_run(caption_text)
    _set_arial(run, size_pt=10, italic=True)

    img_p = _make_blank_para_after(anchor_para)
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img_p.paragraph_format.line_spacing = 1.15
    img_p.paragraph_format.space_before = Pt(6)
    img_p.paragraph_format.space_after = Pt(0)
    run = img_p.add_run()
    run.add_picture(str(png_path), width=Inches(IMAGE_WIDTH_IN))


def main():
    doc = Document(str(DOC))

    inserted = 0
    missing = []
    for spec in LOSS_SPECS:
        anchor = _find_paragraph(doc, spec["anchor"])
        if anchor is None:
            missing.append(spec["anchor"])
            continue
        png = FIG_DIR / spec["png"]
        if not png.exists():
            missing.append(f"(png missing) {spec['png']}")
            continue
        _insert_block(doc, anchor, png, spec["caption"], spec["explanation"])
        inserted += 1
        print(f"  ✓ Equation {spec['eq_num']:>1} inserted after “{spec['anchor'][:45]}…”")

    doc.save(str(DOC))
    print(f"\nInserted {inserted}/{len(LOSS_SPECS)} loss-formula blocks")
    for m in missing:
        print(f"  ✗ {m}")


if __name__ == "__main__":
    main()
